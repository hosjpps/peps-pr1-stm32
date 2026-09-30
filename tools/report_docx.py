"""Оформление отчётов по ГОСТ 7.32: TNR 14, 1.5 интервала, поля 30/15/20/20 мм."""
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

J, C, L, R = (WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER,
              WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT)


class Report:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        st = self.d.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(14)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        pf = st.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = pf.space_after = Pt(0)
        self.fig = self.tab = self.lst = 0
        self._page_numbers()

    def _page_numbers(self):
        sec = self.d.sections[0]
        sec.different_first_page_header_footer = True  # на титульном листе номер не ставится
        p = sec.footer.paragraphs[0]
        p.alignment = C
        run = p.add_run()
        for tag, text in (("begin", None), (None, " PAGE "), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)

    def p(self, text="", align=J, indent=True, bold=False, italic=False, size=None, after=0):
        par = self.d.add_paragraph()
        par.alignment = align
        if indent and align == J:
            par.paragraph_format.first_line_indent = Cm(1.25)
        par.paragraph_format.space_after = Pt(after)
        self._runs(par, text, bold, italic, size)
        return par

    @staticmethod
    def _runs(par, text, bold=False, italic=False, size=None):
        """Поддержка **жирного** и `моноширинного` внутри строки."""
        import re
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = par.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = "Courier New"
                r.font.size = Pt(12)
            else:
                r = par.add_run(part)
                r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)

    def h1(self, text, page_break=False):
        if page_break:
            self.page_break()
        par = self.d.add_paragraph()
        par.alignment = C
        par.paragraph_format.space_before = Pt(12)
        par.paragraph_format.space_after = Pt(12)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(text.upper())
        r.bold = True

    def h2(self, text):
        par = self.d.add_paragraph()
        par.alignment = J
        par.paragraph_format.first_line_indent = Cm(1.25)
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(text)
        r.bold = True

    def bullets(self, items, marker="–"):
        for it in items:
            par = self.d.add_paragraph()
            par.alignment = J
            par.paragraph_format.first_line_indent = Cm(1.25)
            self._runs(par, f"{marker} {it}")

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            par = self.d.add_paragraph()
            par.alignment = J
            par.paragraph_format.first_line_indent = Cm(1.25)
            self._runs(par, f"{i}) {it}")

    def table(self, caption, header, rows, widths=None, size=12):
        self.tab += 1
        cap = self.p(f"Таблица {self.tab} – {caption}", align=L, indent=False)
        cap.paragraph_format.keep_with_next = True
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            self._cell(t.rows[0].cells[i], h, bold=True, size=size, align=C)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                self._cell(cells[i], str(v), size=size)
        if widths:
            t.autofit = False
            for i, w in enumerate(widths):
                t.columns[i].width = Cm(w)  # gridCol — его читают LibreOffice и Word
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.p("", indent=False)
        return self.tab

    def _cell(self, cell, text, bold=False, size=12, align=L):
        par = cell.paragraphs[0]
        par.alignment = align
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.first_line_indent = Cm(0)
        self._runs(par, text, bold=bold, size=size)

    def image(self, path, caption, width_cm=16):
        self.fig += 1
        par = self.d.add_paragraph()
        par.alignment = C
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Cm(width_cm))
        self.p(f"Рисунок {self.fig} – {caption}", align=C, indent=False, after=6)
        return self.fig

    def placeholder(self, caption, what):
        """Место под скриншот, который делается в Proteus/на стенде: красная пометка."""
        self.fig += 1
        par = self.d.add_paragraph()
        par.alignment = C
        r = par.add_run(f"[ВСТАВИТЬ СКРИНШОТ: {what}]")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0, 0)
        self.p(f"Рисунок {self.fig} – {caption}", align=C, indent=False, after=6)
        return self.fig

    def todo(self, text):
        par = self.d.add_paragraph()
        par.alignment = C
        r = par.add_run(f"[{text}]")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0, 0)

    def listing(self, caption, code):
        self.lst += 1
        cap = self.p(f"Листинг {self.lst} – {caption}", align=L, indent=False)
        cap.paragraph_format.keep_with_next = True
        for line in code.rstrip("\n").split("\n"):
            par = self.d.add_paragraph()
            par.alignment = L
            par.paragraph_format.line_spacing = 1.0
            r = par.add_run(line.replace("\t", "    ") or " ")
            r.font.name = "Courier New"
            r.font.size = Pt(10)
        self.p("", indent=False)
        return self.lst

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def title_page(self, work, topic, student, group, teacher="", city_year="Москва 2026"):
        for line in ["МИНОБРНАУКИ РОССИИ",
                     "Федеральное государственное бюджетное образовательное учреждение высшего образования",
                     "«МИРЭА – Российский технологический университет»",
                     "РТУ МИРЭА",
                     "Институт перспективных технологий и индустриального программирования",
                     "Кафедра индустриального программирования"]:
            self.p(line, align=C, indent=False, size=12, bold=line in ("РТУ МИРЭА",))
        for _ in range(3):
            self.p("", indent=False)
        self.p("ОТЧЁТ", align=C, indent=False, bold=True, size=16)
        self.p(f"по {work}", align=C, indent=False)
        self.p("по дисциплине", align=C, indent=False)
        self.p("«Программирование электронных приборов и систем»", align=C, indent=False)
        self.p("", indent=False)
        self.p(f"Тема: «{topic}»", align=C, indent=False, bold=True)
        for _ in range(4):
            self.p("", indent=False)
        for line in [f"Выполнил: студент группы {group}", student, "",
                     f"Проверил: {teacher or '_______________'}"]:
            self.p(line, align=R, indent=False)
        for _ in range(4):
            self.p("", indent=False)
        self.p(city_year, align=C, indent=False)
        self.page_break()

    def save(self, path):
        self.d.save(path)
